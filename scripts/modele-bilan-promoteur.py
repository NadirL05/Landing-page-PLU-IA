"""Génère public/modeles/bilan-promoteur-exemple-fictif.xlsx.

Modèle de bilan promoteur simplifié, proposé en téléchargement sur
/bilan-promoteur. Toutes les valeurs d'hypothèse sont FICTIVES : elles
servent à montrer la mécanique du calcul, pas à donner des références de
marché. Formules volontairement simples (sommes, produits, divisions) pour
rester lisibles dans Excel, LibreOffice ou Numbers.

Usage (hors dépendances du site, aucune installation dans le dépôt) :
    python3 -m venv /tmp/venv-xlsx && /tmp/venv-xlsx/bin/pip install openpyxl
    /tmp/venv-xlsx/bin/python scripts/modele-bilan-promoteur.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

OUT = Path(__file__).resolve().parent.parent / "public" / "modeles" / "bilan-promoteur-exemple-fictif.xlsx"

INPUT = Font(color="1F4FD1", bold=True)
TOTAL = Font(bold=True)
TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, color="555555")
HEAD = PatternFill("solid", fgColor="EDEBE6")
EUR = '#,##0" €"'
EUR_M2 = '#,##0" €/m²"'
M2 = '#,##0" m²"'
PCT = "0.0%"


def header(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = TOTAL
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = HEAD


def line(ws, row, label, value, fmt, note="", kind="formula"):
    ws.cell(row=row, column=1, value=label)
    cell = ws.cell(row=row, column=2, value=value)
    cell.number_format = fmt
    if kind == "input":
        cell.font = INPUT
    elif kind == "total":
        cell.font = TOTAL
        ws.cell(row=row, column=1).font = TOTAL
    ws.cell(row=row, column=3, value="Exemple fictif à remplacer" if kind == "input" else "Calcul")
    ws.cell(row=row, column=4, value=note).font = NOTE


def build():
    wb = Workbook()

    # ---- Lisez-moi -------------------------------------------------------
    info = wb.active
    info.title = "Lisez-moi"
    rows = [
        ("Modèle de bilan promoteur, exemple fictif", TITLE),
        ("Proposé par Fonciable : https://fonciable.fr/bilan-promoteur", None),
        ("", None),
        ("EXEMPLE FICTIF : aucune parcelle réelle, aucune opération réelle, aucun chiffre de client.", TOTAL),
        ("Les valeurs en bleu sont des hypothèses inventées pour illustrer le calcul. Remplacez-les par les vôtres.", None),
        ("Les cellules en noir sont calculées par des formules simples : ne les écrasez pas.", None),
        ("Aucun taux ni prix de ce fichier n'est une référence de marché : ce sont des valeurs d'exemple.", None),
        ("", None),
        ("Onglet « Compte à rebours » : à partir d'une marge cible, calcule la charge foncière admissible", None),
        ("et le prix maximal du terrain hors frais d'acquisition.", None),
        ("Onglet « Bilan direct » : à partir d'un prix de terrain négocié, calcule la marge obtenue.", None),
        ("", None),
        ("Simplifications assumées : montants hors taxes ; frais financiers forfaitaires en % du chiffre", None),
        ("d'affaires (sans échéancier) ; taxe d'aménagement et participations saisies en montant ;", None),
        ("aucune fiscalité propre à votre structure.", None),
        ("", None),
        ("Ce fichier est une aide au raisonnement. Ce n'est ni un certificat d'urbanisme, ni un avis", None),
        ("juridique, fiscal ou financier. La surface de plancher réellement constructible se confirme", None),
        ("auprès du service urbanisme de la commune ; les coûts, par des devis.", None),
        ("", None),
        ("Postes et formules : ceux de la page https://fonciable.fr/bilan-promoteur (méthode du compte à rebours).", None),
        ("Dans l'application Fonciable, le bilan part des données publiques de la parcelle et d'hypothèses", None),
        ("de départ que vous ajustez : https://app.fonciable.fr", None),
    ]
    for i, (text, font) in enumerate(rows, start=1):
        cell = info.cell(row=i, column=1, value=text)
        if font is not None:
            cell.font = font
    info.column_dimensions["A"].width = 110

    # ---- Compte à rebours -----------------------------------------------
    ws = wb.create_sheet("Compte à rebours")
    ws["A1"] = "Bilan promoteur à rebours : charge foncière admissible (exemple fictif)"
    ws["A1"].font = TITLE
    for col, text in zip("ABCD", ["Poste", "Valeur", "Nature", "Commentaire"]):
        ws[f"{col}3"] = text
        ws[f"{col}3"].font = TOTAL

    # Postes repris de la page /bilan-promoteur (content/bilan-promoteur.ts,
    # bloc « Les postes d'un bilan promoteur ») et formules de son bloc
    # « La méthode du compte à rebours ». Ne pas ajouter de poste ni de taux
    # qui n'y figurent pas. Les numéros de ligne sont calculés : une ligne
    # ajoutée ne casse aucune référence.
    r = {}  # clé du poste -> référence de cellule (ex. "B13")
    row_ref = [4]  # prochaine ligne libre

    def put(key, label, value, fmt, note="", kind="formula"):
        n = row_ref[0]
        line(ws, n, label, value.format(**r) if isinstance(value, str) else value, fmt, note, kind)
        r[key] = f"B{n}"
        row_ref[0] += 1

    def section(title):
        header(ws, row_ref[0], title)
        row_ref[0] += 1

    def gap():
        row_ref[0] += 1

    section("Programme")
    put("sdp", "Surface de plancher (SDP)", 1240, M2, "Issue des règles du PLU : à confirmer en mairie.", "input")
    put("ratio", "Rapport surface vendable / SDP", 0.9, PCT, "Part de la SDP réellement vendue.", "input")
    put("sv", "Surface vendable", "={sdp}*{ratio}", M2)
    gap()

    section("Recettes (HT)")
    put("pv", "Prix de vente moyen HT", 4200, EUR_M2, "Par m² vendable.", "input")
    put("ventes", "Ventes de logements et de locaux HT", "={sv}*{pv}", EUR)
    put("autres", "Places de stationnement et autres recettes HT", 0, EUR, "Montant total ; 0 si aucune.", "input")
    put("ca", "Chiffre d'affaires HT", "={ventes}+{autres}", EUR, "", "total")
    gap()

    section("Dépenses hors foncier (HT)")
    put("cc", "Coût de construction HT", 1650, EUR_M2, "Par m² de SDP.", "input")
    put("trav", "Travaux de construction", "={sdp}*{cc}", EUR)
    put("ssol", "Sous-sols et adaptation au terrain (pente, fondations)", 0, EUR, "Montant total ; 0 si sans objet.", "input")
    put("travt", "Total des travaux", "={trav}+{ssol}", EUR)
    put("honp", "Honoraires techniques (% des travaux)", 0.10, PCT, "Architecte, bureaux d'études, géomètre, contrôle…", "input")
    put("hon", "Honoraires techniques", "={travt}*{honp}", EUR)
    put("taxe", "Taxe d'aménagement, archéologie préventive, participations", 80000, EUR, "Montant à calculer pour la commune.", "input")
    put("assp", "Assurances (% des travaux)", 0.02, PCT, "Dommages-ouvrage notamment.", "input")
    put("ass", "Assurances", "={travt}*{assp}", EUR)
    put("alep", "Aléas (% des travaux)", 0.03, PCT, "Provision pour imprévus.", "input")
    put("ale", "Aléas", "={travt}*{alep}", EUR)
    put("comp", "Frais de commercialisation (% du CA)", 0.04, PCT, "Honoraires de vente, publicité…", "input")
    put("com", "Frais de commercialisation", "={ca}*{comp}", EUR)
    put("finp", "Frais financiers (% du CA)", 0.025, PCT, "Forfait simplifié, sans échéancier.", "input")
    put("fin", "Frais financiers", "={ca}*{finp}", EUR)
    put("dep", "Total des dépenses hors foncier", "={travt}+{hon}+{taxe}+{ass}+{ale}+{com}+{fin}", EUR, "", "total")
    gap()

    section("Marge et foncier")
    put("margep", "Marge cible (% du CA HT)", 0.12, PCT, "Objectif de l'opérateur.", "input")
    put("marge", "Marge cible", "={ca}*{margep}", EUR)
    put("cfa", "Charge foncière admissible (frais d'acquisition inclus)", "={ca}-{marge}-{dep}", EUR,
        "CA HT − marge cible − dépenses hors foncier.", "total")
    put("fraisp", "Frais d'acquisition (% du prix du terrain)", 0.075, PCT, "Notaire, droits : à confirmer.", "input")
    put("pmax", "Prix maximal du terrain hors frais", "={cfa}/(1+{fraisp})", EUR, "Charge foncière ÷ (1 + frais). Base de l'offre.", "total")
    put("inc", "Incidence foncière par m² de SDP", "={cfa}/{sdp}", EUR_M2, "Charge foncière ÷ SDP.")
    put("part", "Part du foncier dans le CA HT", "={cfa}/{ca}", PCT)

    # ---- Bilan direct ----------------------------------------------------
    wd = wb.create_sheet("Bilan direct")
    wd["A1"] = "Bilan promoteur direct : marge pour un prix de terrain donné (exemple fictif)"
    wd["A1"].font = TITLE
    for col, text in zip("ABCD", ["Poste", "Valeur", "Nature", "Commentaire"]):
        wd[f"{col}3"] = text
        wd[f"{col}3"].font = TOTAL
    src = "'Compte à rebours'!"
    header(wd, 4, "Repris de l'onglet « Compte à rebours »")
    line(wd, 5, "Chiffre d'affaires HT", f"={src}{r['ca']}", EUR)
    line(wd, 6, "Total des dépenses hors foncier", f"={src}{r['dep']}", EUR)
    line(wd, 7, "Frais d'acquisition (% du prix du terrain)", f"={src}{r['fraisp']}", PCT)
    header(wd, 9, "Foncier négocié")
    line(wd, 10, "Prix du terrain négocié", 1200000, EUR, "", "input")
    line(wd, 11, "Frais d'acquisition", "=B10*B7", EUR)
    line(wd, 12, "Charge foncière", "=B10+B11", EUR, "", "total")
    header(wd, 14, "Résultat")
    line(wd, 15, "Prix de revient total", "=B12+B6", EUR, "", "total")
    line(wd, 16, "Marge", "=B5-B15", EUR, "", "total")
    line(wd, 17, "Marge en % du CA HT", "=B16/B5", PCT, "", "total")

    for sheet in (ws, wd):
        sheet.column_dimensions["A"].width = 58
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 48
        sheet.freeze_panes = "A4"
        for row in sheet.iter_rows(min_row=4, min_col=2, max_col=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="right")

    wb.properties.title = "Modèle de bilan promoteur (exemple fictif)"
    wb.properties.creator = "Fonciable"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
